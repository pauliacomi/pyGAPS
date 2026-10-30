"""
Parse to and from a Excel format for isotherms.

Isotherms are written as ``.xlsx`` files using ``openpyxl``. Both ``.xlsx``
and legacy ``.xls`` files can be read, the latter through ``xlrd``.

The _parser_version variable documents any changes to the format,
and is used to check for any deprecations.

"""

import ast
import math
import os
import pathlib
import posixpath
import tempfile
import warnings
import zipfile
from xml.etree import ElementTree

import pandas

from pygaps import logger
from pygaps.core.baseisotherm import BaseIsotherm
from pygaps.core.modelisotherm import ModelIsotherm
from pygaps.core.pointisotherm import PointIsotherm
from pygaps.modelling import model_from_dict
from pygaps.utilities.exceptions import ParsingError

_parser_version = "3.0"

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

_META_DICT = {
    'material': {
        'text': ('Material name',),
        'name': 'material',
        'row': 0,
        'column': 0,
    },
    'temperature': {
        'text': ('Experiment temperature (K)',),
        'name': 'temperature',
        'row': 1,
        'column': 0,
    },
    'adsorbate': {
        'text': ('Adsorbate used',),
        'name': 'adsorbate',
        'row': 2,
        'column': 0,
    },
    'pressure_mode': {
        'text': ('Pressure mode',),
        'name': 'pressure_mode',
        'row': 3,
        'column': 0,
    },
    'pressure_unit': {
        'text': ('Pressure unit',),
        'name': 'pressure_unit',
        'row': 4,
        'column': 0,
    },
    'loading_basis': {
        'text': ('Loading basis',),
        'name': 'loading_basis',
        'row': 5,
        'column': 0,
    },
    'loading_unit': {
        'text': ('Loading unit',),
        'name': 'loading_unit',
        'row': 6,
        'column': 0,
    },
    'material_basis': {
        'text': ('Material basis',),
        'name': 'material_basis',
        'row': 7,
        'column': 0,
    },
    'material_unit': {
        'text': ('Material unit',),
        'name': 'material_unit',
        'row': 8,
        'column': 0,
    },
    'isotherm_data': {
        'text': ('Isotherm type',),
        'name': 'isotherm_data',
        'row': 9,
        'column': 0,
    },
}


def _cell_value(val):
    """Convert numpy scalars to Python types openpyxl can store."""
    if hasattr(val, 'item'):
        return val.item()
    return val


def _preserve_float_precision(path, float_values):
    """Write floats using their shortest round-trippable decimal representation."""
    if not float_values:
        return

    with zipfile.ZipFile(path, 'r') as source:
        workbook = ElementTree.fromstring(source.read('xl/workbook.xml'))
        relationships = ElementTree.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        targets = {
            relationship.attrib['Id']: relationship.attrib['Target']
            for relationship in relationships.findall(f"{{{_PACKAGE_REL_NS}}}Relationship")
        }
        worksheets = {
            sheet.attrib['name']: posixpath.normpath(
                sheet_target if sheet_target.startswith('/') else posixpath.join('xl', sheet_target)
            ).lstrip('/')
            for sheet in workbook.findall(f".//{{{_MAIN_NS}}}sheet")
            if (sheet_target := targets.get(sheet.attrib[f"{{{_DOCUMENT_REL_NS}}}id"]))
        }

        updated_sheets = {}
        for sheet_name, cells in float_values.items():
            worksheet_path = worksheets.get(sheet_name)
            if worksheet_path is None:
                raise ParsingError(f"Could not find worksheet '{sheet_name}' in saved Excel file")

            worksheet = ElementTree.fromstring(source.read(worksheet_path))
            found_cells = set()
            for cell in worksheet.iter(f"{{{_MAIN_NS}}}c"):
                coordinate = cell.attrib.get('r')
                if coordinate not in cells:
                    continue
                value = cell.find(f"{{{_MAIN_NS}}}v")
                if value is None:
                    raise ParsingError(f"Could not update numeric value in cell {sheet_name}!{coordinate}")
                value.text = cells[coordinate]
                found_cells.add(coordinate)

            if found_cells != cells.keys():
                missing = ', '.join(sorted(cells.keys() - found_cells))
                raise ParsingError(f"Could not find numeric cells in worksheet '{sheet_name}': {missing}")

            updated_sheets[worksheet_path] = ElementTree.tostring(worksheet, encoding='utf-8', xml_declaration=True)

    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=path.suffix, delete=False) as temporary_file:
        temporary_path = pathlib.Path(temporary_file.name)

    try:
        with zipfile.ZipFile(path, 'r') as source, zipfile.ZipFile(temporary_path, 'w') as destination:
            destination.comment = source.comment
            for entry in source.infolist():
                data = updated_sheets.get(entry.filename)
                if data is None:
                    data = source.read(entry.filename)
                destination.writestr(entry, data)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def isotherm_to_xl(isotherm, path):
    """
    Save an isotherm to an excel (.xlsx) file.

    Numeric values are stored with enough precision to preserve their original
    floating-point values when read back by pyGAPS.

    Parameters
    ----------
    isotherm : Isotherm
        Isotherm to be written to excel.
    path : str
        Path to the file to be written. Legacy ``.xls`` files cannot be
        written, so a ``.xls`` extension is replaced with ``.xlsx``.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.styles import PatternFill

    path = pathlib.Path(path)
    if path.suffix.lower() == '.xls':
        new_path = path.with_suffix('.xlsx')
        warnings.warn(
            f"Legacy .xls files can no longer be written. Saving to '{new_path}' instead.",
            stacklevel=2,
        )
        path = new_path

    # create a new workbook and select first sheet
    wb = Workbook()
    sht = wb.active
    sht.title = 'data'
    float_values = {}

    def write(row, col, val, styled=False):
        value = _cell_value(val)
        cell = sht.cell(row=row + 1, column=col + 1, value=value)
        if isinstance(value, float) and math.isfinite(value):
            float_values.setdefault(sht.title, {})[cell.coordinate] = repr(value)
        if styled:
            cell.alignment = Alignment(horizontal='left')
            cell.fill = PatternFill('solid', fgColor='C0C0C0')

    # get the required dictionaries
    iso_dict = isotherm.to_dict()
    iso_dict['file_version'] = _parser_version  # version

    # Parse material
    material = iso_dict['material']
    if isinstance(material, dict):
        iso_dict['material'] = material.pop('name')
        iso_dict.update({f"_material_{key}": val for key, val in material.items()})

    # Add the required named properties
    for field in _META_DICT.values():
        val = iso_dict.pop(field['name'], None)
        write(field['row'], field['column'], field['text'][0], styled=True)
        if val:
            write(field['row'], field['column'] + 1, val, styled=True)

    # Get the isotherm type header
    type_row = _META_DICT['isotherm_data']['row']
    type_col = _META_DICT['isotherm_data']['column']

    sht.column_dimensions['A'].width = 25
    sht.column_dimensions['B'].width = 25

    if isinstance(isotherm, PointIsotherm):
        # Write the type
        write(type_row, type_col + 1, 'data', styled=True)

        # Get the data row
        data_row = type_row + 1

        # We get data and replace adsorption terminology
        data = isotherm.data_raw.copy()
        data['branch'] = data['branch'].replace(0, 'ads').replace(1, 'des')

        columns = [isotherm.pressure_key, isotherm.loading_key, 'branch'] + isotherm.other_keys

        # Write data types
        for col_index, heading in enumerate(isotherm.other_keys):
            write(data_row - 1, col_index + 3, data[heading].dtype.name)

        # Write all data
        for col_index, heading in enumerate(columns):
            write(data_row, col_index, heading)
            for row_index, datapoint in enumerate(data[heading]):
                write(data_row + row_index + 1, col_index, datapoint)

    elif isinstance(isotherm, ModelIsotherm):
        # Write the type
        write(type_row, type_col + 1, 'model', styled=True)

        # Get the model row
        model_row = type_row

        # Generate the headings (i know it's a lame way to do it)
        write(model_row + 1, 0, 'Model name')
        write(model_row + 1, 1, isotherm.model.name)
        write(model_row + 2, 0, 'RMSE')
        write(model_row + 2, 1, isotherm.model.rmse)
        write(model_row + 3, 0, 'Pressure range')
        write(model_row + 3, 1, f'({isotherm.model.pressure_range[0]}, {isotherm.model.pressure_range[1]})')
        write(model_row + 4, 0, 'Loading range')
        write(model_row + 4, 1, f'({isotherm.model.loading_range[0]}, {isotherm.model.loading_range[1]})')
        write(model_row + 5, 0, 'Model parameters')
        model_row = model_row + 5
        for row_index, param in enumerate(isotherm.model.params):
            write(model_row + row_index + 1, 0, param)
            write(model_row + row_index + 1, 1, isotherm.model.params[param])

    elif isinstance(isotherm, BaseIsotherm):
        # Write the type
        write(type_row, type_col + 1, 'metadata', styled=True)

    else:
        raise ParsingError("Could not determine isotherm type")

    # Now add the other keys
    sht = wb.create_sheet('otherdata')
    for row, prop in enumerate(iso_dict):
        write(row, 0, prop)
        write(row, 1, iso_dict[prop])

    wb.save(path)
    _preserve_float_precision(path, float_values)


class _Sheet:
    """Read-only, 0-indexed view of a worksheet from either xlrd or openpyxl.

    Empty cells are returned as ``None`` and booleans as ``bool``.
    """

    def __init__(self, values):
        self._values = values
        self.nrows = len(values)
        self.ncols = max((len(row) for row in values), default=0)

    def value(self, row, col):
        try:
            return self._values[row][col]
        except IndexError:
            return None


def _read_workbook(path):
    """Read all sheets of an excel file into a ``{name: _Sheet}`` dict."""
    path = pathlib.Path(path)
    sheets = {}

    if path.suffix.lower() == '.xls':
        import xlrd

        wb = xlrd.open_workbook(path)
        for sht in wb.sheets():
            rows = []
            for r in range(sht.nrows):
                row = []
                for cell in sht.row(r):
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        row.append(bool(cell.value))
                    else:
                        row.append(cell.value)
                rows.append(row)
            sheets[sht.name] = _Sheet(rows)
    else:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)

        # Excel has a single number type: return floats like xlrd does
        def normalise(val):
            if isinstance(val, int) and not isinstance(val, bool):
                return float(val)
            return val

        try:
            for sht in wb.worksheets:
                sheets[sht.title] = _Sheet([[normalise(v) for v in row] for row in sht.iter_rows(values_only=True)])
        finally:
            wb.close()

    return sheets


def isotherm_from_xl(path, **isotherm_parameters):
    """
    Load an isotherm from a pyGAPS Excel file (.xlsx or legacy .xls).

    Parameters
    ----------
    path : str
        Path to the file to be read.
    isotherm_parameters :
        Any other options to be overridden in the isotherm creation.

    Returns
    -------
    Isotherm
        The isotherm contained in the excel file.

    """
    # isotherm type (point/model)
    isotype = 0

    raw_dict = {}

    # Get excel workbook and sheet
    sheets = _read_workbook(path)
    if 'data' in sheets:
        sht = sheets['data']
    else:
        sht = next(iter(sheets.values()))

    def empty(val):
        return val is None or val == ''

    # read the main isotherm parameters
    for field in _META_DICT.values():
        raw_dict[field['name']] = sht.value(field['row'], field['column'] + 1)

    # find data/model limits
    type_row = _META_DICT['isotherm_data']['row']
    isotherm_type = str(sht.value(type_row, 1) or '').lower()

    if isotherm_type.startswith('data'):
        # Store isotherm type
        isotype = 1

        header_row = type_row + 1
        start_row = header_row + 1
        final_row = start_row

        while final_row < sht.nrows:
            if empty(sht.value(final_row, 0)):
                break
            final_row += 1

        # read the data in
        header_col = 0
        headers = []
        dtypes = {}
        experiment_data = {}
        while header_col < sht.ncols:
            header = sht.value(header_row, header_col)
            if empty(header):
                break
            # read header, data, and dtype
            headers.append(header)
            experiment_data[header] = [sht.value(i, header_col) for i in range(start_row, final_row)]
            # read the data type
            # only read the data type if it is not a pressure, loading or branch
            if header_col > 2:
                dtype = sht.value(header_row - 1, header_col)
                if dtype != '':
                    dtypes[header] = dtype
            header_col += 1
        data = pandas.DataFrame(experiment_data)
        data = data.astype(dtypes)

        raw_dict['pressure_key'] = headers[0]
        raw_dict['loading_key'] = headers[1]

        # process isotherm branches if they exist
        if 'branch' in data.columns:
            data['branch'] = data['branch'].apply(lambda x: 0 if x == 'ads' else 1)
        else:
            raw_dict['branch'] = 'guess'

    if isotherm_type.startswith('model'):
        # Store isotherm type
        isotype = 2
        model = {
            "name": sht.value(type_row + 1, 1),
            "rmse": sht.value(type_row + 2, 1),
            "pressure_range": ast.literal_eval(sht.value(type_row + 3, 1)),
            "loading_range": ast.literal_eval(sht.value(type_row + 4, 1)),
            "parameters": {},
        }

        final_row = type_row + 6

        while final_row < sht.nrows:
            point = sht.value(final_row, 0)
            if empty(point):
                break
            model["parameters"][point] = sht.value(final_row, 1)
            final_row += 1

    # read the secondary isotherm parameters
    if 'otherdata' in sheets:
        sht = sheets['otherdata']
        row_index = 0
        while row_index < sht.nrows:
            name = sht.value(row_index, 0)
            if empty(name):
                break
            val = sht.value(row_index, 1)
            raw_dict[name] = val
            row_index += 1

    # Put data in order
    # version check
    version = raw_dict.pop("file_version", None)
    if not version or float(version) < float(_parser_version):
        logger.warning(
            f"The file version is {version} while the parser uses version {_parser_version}. "
            "Strange things might happen, so double check your data."
        )
    raw_dict.pop('isotherm_data')  # remove useless field
    raw_dict.pop('iso_id', None)  # make sure id is not passed

    # check if material needs parsing
    material = {}
    for key, val in raw_dict.items():
        if key.startswith("_material_"):
            material[key.replace("_material_", "")] = val
    if material:
        for key in material:
            raw_dict.pop("_material_" + key)
        material['name'] = raw_dict['material']
        raw_dict['material'] = material

    # Update dictionary with any user parameters
    raw_dict.update(isotherm_parameters)

    if isotype == 1:
        return PointIsotherm(isotherm_data=data, **raw_dict)

    if isotype == 2:
        return ModelIsotherm(model=model_from_dict(model), **raw_dict)

    return BaseIsotherm(**raw_dict)
